"""Regression tests for organization contract ingest and validation."""

from __future__ import annotations

import json

import pytest
import yaml
from docx import Document

from irb_harness.application.contract_compiler import compile_contract_from_documents
from irb_harness.domain.contracts import ContractValidationError, OrganizationContract
from irb_harness.infrastructure.contract_loader import contract_sha256, load_contract
from irb_harness.infrastructure.document_ingest import ingest_document


def test_bundled_default_contract_is_kmuh_and_valid():
    contract = load_contract()

    assert contract.organization_id == "kmuh"
    assert contract.contract_id == "kmuh-irb"
    assert {site.site_id for site in contract.websites} == {
        "kmuh_ctc_directory",
        "kmuh_eirb",
    }
    assert len(contract.form_sets) >= 8
    assert len(contract.source_documents) >= 11
    assert len(contract.requirements) == 9
    assert all(transition.evidence_refs for transition in contract.workflow_transitions)
    assert all(requirement.evidence_refs for requirement in contract.requirements)
    assert set(contract.summary()["workflow_transitions_needing_evidence"]) == {
        transition.event for transition in contract.workflow_transitions
    }
    assert len(contract_sha256(contract)) == 64


def test_mapped_portal_requirement_requires_matching_reviewed_binding():
    mapping = {
        "schema_version": "1.0",
        "contract_id": "example",
        "organization": {"id": "example", "name": "Example"},
        "source_documents": [],
        "workflow": {
            "states": ["draft", "submitted"],
            "transitions": [
                {
                    "event": "submit",
                    "from": "draft",
                    "to": "submitted",
                    "action_risk": "submit",
                    "human_confirmation": True,
                }
            ],
        },
        "requirements": [
            {
                "requirement_id": "portal_title",
                "title": "Title",
                "kind": "portal_field",
                "status": "needs_evidence",
                "data_path": "study.title",
                "website_site_id": "portal",
                "page_mapping_sha256": "a" * 64,
                "control_id": "b" * 64,
            }
        ],
        "form_sets": [],
        "websites": [
            {
                "site_id": "portal",
                "title": "Portal",
                "start_url": "https://example.org/",
                "allowed_hosts": ["example.org"],
                "login": {"mode": "human"},
                "action_policy": {"read_actions": "reviewed_only"},
            }
        ],
    }

    with pytest.raises(ContractValidationError, match="requires a website binding"):
        OrganizationContract.from_mapping(mapping)


def test_submit_transition_requires_human_confirmation():
    invalid = {
        "schema_version": "1.0",
        "contract_id": "example",
        "organization": {"id": "example", "name": "Example"},
        "source_documents": [],
        "workflow": {
            "states": ["draft", "submitted"],
            "transitions": [
                {
                    "event": "submit",
                    "from": "draft",
                    "to": "submitted",
                    "action_risk": "submit",
                    "human_confirmation": False,
                }
            ],
        },
        "form_sets": [],
        "websites": [],
    }

    with pytest.raises(ContractValidationError, match="human_confirmation"):
        OrganizationContract.from_mapping(invalid)


def test_verified_workflow_evidence_requires_spans_and_hashed_source():
    invalid = {
        "schema_version": "1.0",
        "contract_id": "example",
        "organization": {"id": "example", "name": "Example"},
        "source_documents": [
            {
                "source_id": "sop",
                "title": "SOP",
                "uri": "https://example.org/sop.pdf",
                "media_type": "application/pdf",
            }
        ],
        "workflow": {
            "states": ["draft", "submitted"],
            "transitions": [
                {
                    "event": "submit",
                    "from": "draft",
                    "to": "submitted",
                    "action_risk": "submit",
                    "human_confirmation": True,
                    "evidence_refs": [
                        {
                            "source_document_id": "sop",
                            "locator_status": "verified",
                            "span_ids": [],
                        }
                    ],
                }
            ],
        },
        "form_sets": [],
        "websites": [],
    }

    with pytest.raises(ContractValidationError, match="at least one span_id"):
        OrganizationContract.from_mapping(invalid)

    invalid["workflow"]["transitions"][0]["evidence_refs"][0]["span_ids"] = ["sop:L1"]
    with pytest.raises(ContractValidationError, match="unhashed source"):
        OrganizationContract.from_mapping(invalid)

    invalid["source_documents"][0]["sha256"] = "c" * 64
    with pytest.raises(ContractValidationError, match="matching evidence binding"):
        OrganizationContract.from_mapping(invalid)


def test_organization_id_cannot_escape_workspace_paths():
    invalid = {
        "schema_version": "1.0",
        "contract_id": "escape",
        "organization": {"id": "../../escape", "name": "Unsafe"},
        "source_documents": [],
        "workflow": {"states": [], "transitions": []},
        "form_sets": [],
        "websites": [],
    }

    with pytest.raises(ContractValidationError, match="filesystem-safe"):
        OrganizationContract.from_mapping(invalid)


def test_docx_ingest_preserves_hashes_and_line_locators(tmp_path):
    source = tmp_path / "KMUH 新案表.docx"
    document = Document()
    document.add_paragraph("計畫名稱：測試研究")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "主持人"
    table.cell(0, 1).text = "王大明"
    document.save(source)

    ingested = ingest_document(source, institution_id="kmuh")

    assert ingested.source_id.startswith("kmuh:")
    assert len(ingested.byte_sha256) == 64
    assert any(span.context == "計畫名稱：測試研究" for span in ingested.spans)
    assert all(span.char_end >= span.char_start for span in ingested.spans)


def test_compile_contract_creates_separate_evidence_index_without_overwrite(tmp_path):
    source = tmp_path / "instructions.md"
    source.write_text("新案送審\n人工確認後送出\n", encoding="utf-8")
    target = tmp_path / "organization.yml"

    contract_path, evidence_path = compile_contract_from_documents(
        "example",
        "Example IRB",
        [source],
        target,
    )

    contract = yaml.safe_load(contract_path.read_text(encoding="utf-8"))
    evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
    assert contract["organization"]["id"] == "example"
    assert contract["status"] == "draft"
    assert contract["compiler"]["rule_inference"] == "not_performed"
    assert evidence["documents"][0]["spans"][0]["line_start"] == 1

    with pytest.raises(FileExistsError, match="refusing to overwrite"):
        compile_contract_from_documents("example", "Example IRB", [source], target)
