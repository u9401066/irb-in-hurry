"""Narrow document ingest adapters with citation-ready evidence spans."""

from __future__ import annotations

import hashlib
import mimetypes
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable


@dataclass(frozen=True)
class EvidenceSpan:
    span_id: str
    line_start: int
    line_end: int
    char_start: int
    char_end: int
    text_sha256: str
    context: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "span_id": self.span_id,
            "locator_version": "1",
            "line_start": self.line_start,
            "line_end": self.line_end,
            "char_start": self.char_start,
            "char_end": self.char_end,
            "text_sha256": self.text_sha256,
            "context": self.context,
        }


@dataclass(frozen=True)
class IngestedDocument:
    source_id: str
    path: Path
    media_type: str
    byte_sha256: str
    byte_size: int
    text: str
    spans: tuple[EvidenceSpan, ...]

    def contract_source(self) -> dict[str, Any]:
        return {
            "source_id": self.source_id,
            "title": self.path.name,
            "uri": self.path.resolve().as_uri(),
            "media_type": self.media_type,
            "sha256": self.byte_sha256,
            "evidence_status": "ingested_needs_rule_mapping",
        }

    def evidence_record(self) -> dict[str, Any]:
        return {
            "source_id": self.source_id,
            "uri": self.path.resolve().as_uri(),
            "byte_sha256": self.byte_sha256,
            "byte_size": self.byte_size,
            "media_type": self.media_type,
            "text_sha256": hashlib.sha256(self.text.encode("utf-8")).hexdigest(),
            "spans": [span.as_dict() for span in self.spans],
        }


def ingest_document(path: str | Path, *, institution_id: str) -> IngestedDocument:
    """Extract text and stable locators without modifying the source document."""
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"document not found: {source}")
    raw = source.read_bytes()
    byte_sha256 = hashlib.sha256(raw).hexdigest()
    media_type = _media_type(source)
    text = _extract_text(source, media_type)
    normalized = _normalize_text(text)
    source_id = f"{institution_id}:{_slug(source.stem)}:{byte_sha256[:12]}"
    spans = tuple(_make_spans(source_id, normalized))
    return IngestedDocument(
        source_id=source_id,
        path=source,
        media_type=media_type,
        byte_sha256=byte_sha256,
        byte_size=len(raw),
        text=normalized,
        spans=spans,
    )


def _extract_text(path: Path, media_type: str) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in {".html", ".htm"} or media_type == "text/html":
        return _extract_html(path)
    if suffix in {".txt", ".md", ".csv", ".yml", ".yaml", ".json"}:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"unsupported document type: {media_type} ({path.suffix})")


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    lines: list[str] = []
    _append_paragraphs(lines, document.paragraphs)
    for table_index, table in enumerate(document.tables, start=1):
        lines.append(f"[TABLE {table_index}]")
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    for section_index, section in enumerate(document.sections, start=1):
        lines.append(f"[HEADER {section_index}]")
        _append_paragraphs(lines, section.header.paragraphs)
        lines.append(f"[FOOTER {section_index}]")
        _append_paragraphs(lines, section.footer.paragraphs)
    return "\n".join(lines)


def _append_paragraphs(lines: list[str], paragraphs: Iterable[Any]) -> None:
    for paragraph in paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return "\n\f\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_html(path: Path) -> str:
    parser = _VisibleTextParser()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    parser.close()
    return "\n".join(parser.lines)


class _VisibleTextParser(HTMLParser):
    """Small dependency-free extractor that omits script/style/template bodies."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.lines: list[str] = []
        self._hidden_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag.lower() in {"script", "style", "template", "noscript"}:
            self._hidden_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "template", "noscript"}:
            self._hidden_depth = max(0, self._hidden_depth - 1)

    def handle_data(self, data: str) -> None:
        if self._hidden_depth:
            return
        normalized = re.sub(r"\s+", " ", data).strip()
        if normalized:
            self.lines.append(normalized)


def _make_spans(source_id: str, text: str) -> Iterable[EvidenceSpan]:
    offset = 0
    for line_number, line in enumerate(text.splitlines(keepends=True), start=1):
        content = line.rstrip("\r\n")
        start = offset
        end = start + len(content)
        offset += len(line)
        if not content.strip():
            continue
        digest = hashlib.sha256(content.encode("utf-8")).hexdigest()
        yield EvidenceSpan(
            span_id=f"{source_id}:L{line_number}",
            line_start=line_number,
            line_end=line_number,
            char_start=start,
            char_end=end,
            text_sha256=digest,
            context=content[:240],
        )


def _normalize_text(text: str) -> str:
    return "\n".join(
        line.rstrip()
        for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    )


def _media_type(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if path.suffix.lower() in {".yml", ".yaml"}:
        return "application/yaml"
    return mimetypes.guess_type(path.name)[0] or "application/octet-stream"


def _slug(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug or "document"
